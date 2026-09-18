from collections.abc import Callable
import io

from docx import Document
from fastapi.testclient import TestClient


def cv_file() -> tuple[str, bytes, str]:
    document = Document()
    document.add_paragraph("Mouid Test")
    document.add_paragraph(
        "Built Python and FastAPI REST APIs with PostgreSQL and Docker. "
        "Worked in a collaborative team and tested production services."
    )
    content = io.BytesIO()
    document.save(content)
    return (
        "cv.docx",
        content.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def test_cv_analysis_and_application_toolkit_workflow(
    client: TestClient,
    register: Callable[[str, str], dict[str, str]],
    create_application: Callable[[dict[str, str], dict[str, object] | None], dict[str, object]],
) -> None:
    owner = register("Mouid Test", "owner@example.com")
    outsider = register("Other User", "other@example.com")
    application = create_application(owner)
    application_id = application["id"]

    uploaded = client.post(
        f"/api/applications/{application_id}/documents/cv",
        headers=owner,
        files={"file": cv_file()},
    )
    assert uploaded.status_code == 201
    assert uploaded.json()["extracted_characters"] > 50

    analysis = client.post(f"/api/applications/{application_id}/analysis", headers=owner)
    assert analysis.status_code == 200
    assert analysis.json()["is_stale"] is False
    assert "Python" in analysis.json()["matching_skills"]

    generated = client.post(f"/api/applications/{application_id}/application-kit", headers=owner)
    assert generated.status_code == 200
    kit = generated.json()
    assert len(kit["interview_questions"]) == 6
    assert len(kit["questions_to_ask"]) == 4
    assert "Example Labs" in kit["cover_letter"]
    assert kit["is_stale"] is False

    saved = client.get(f"/api/applications/{application_id}/application-kit", headers=owner)
    assert saved.status_code == 200
    assert saved.json()["id"] == kit["id"]
    assert client.get(f"/api/applications/{application_id}/application-kit", headers=outsider).status_code == 404

    changed = client.patch(
        f"/api/applications/{application_id}",
        headers=owner,
        json={"job_description": "Build Python services with Kubernetes and AWS."},
    )
    assert changed.status_code == 200
    assert client.get(f"/api/applications/{application_id}/analysis", headers=owner).json()["is_stale"] is True
    assert client.get(f"/api/applications/{application_id}/application-kit", headers=owner).json()["is_stale"] is True
    assert client.post(f"/api/applications/{application_id}/application-kit", headers=owner).status_code == 422


def test_analysis_requires_cv_and_job_description(
    client: TestClient,
    register: Callable[[str, str], dict[str, str]],
    create_application: Callable[[dict[str, str], dict[str, object] | None], dict[str, object]],
) -> None:
    headers = register("Test User", "test@example.com")
    application = create_application(headers, {"job_description": None})
    application_id = application["id"]

    missing_description = client.post(f"/api/applications/{application_id}/analysis", headers=headers)
    assert missing_description.status_code == 422

    client.patch(
        f"/api/applications/{application_id}",
        headers=headers,
        json={"job_description": "Python and FastAPI"},
    )
    missing_cv = client.post(f"/api/applications/{application_id}/analysis", headers=headers)
    assert missing_cv.status_code == 422
